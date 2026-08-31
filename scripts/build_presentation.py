"""Build the submission PowerPoint from verified project results."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "presentation" / "CORES_MDE_presentation.pptx"
NAVY = RGBColor(24, 45, 78)
BLUE = RGBColor(68, 114, 196)
ORANGE = RGBColor(237, 125, 49)
DARK = RGBColor(35, 35, 35)
LIGHT = RGBColor(245, 247, 250)


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    title_box = slide.shapes.add_textbox(Inches(0.65), Inches(0.35), Inches(12), Inches(0.65))
    paragraph = title_box.text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.size = Pt(28)
    paragraph.font.bold = True
    paragraph.font.color.rgb = NAVY
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.68), Inches(1.0), Inches(11.8), Inches(0.35))
        sub = sub_box.text_frame.paragraphs[0]
        sub.text = subtitle
        sub.font.size = Pt(12)
        sub.font.color.rgb = RGBColor(90, 90, 90)


def add_bullets(slide, bullets: list[str], *, left=0.8, top=1.35, width=5.5, height=5.5) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    for index, text in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.level = 0
        paragraph.font.size = Pt(20)
        paragraph.font.color.rgb = DARK
        paragraph.space_after = Pt(14)


def add_picture(slide, filename: str, *, left=6.3, top=1.25, width=6.5) -> None:
    slide.shapes.add_picture(str(ROOT / "figures" / filename), Inches(left), Inches(top), width=Inches(width))


def add_footer(slide, number: int) -> None:
    box = slide.shapes.add_textbox(Inches(11.8), Inches(7.05), Inches(0.8), Inches(0.25))
    p = box.text_frame.paragraphs[0]
    p.text = str(number)
    p.alignment = PP_ALIGN.RIGHT
    p.font.size = Pt(10)
    p.font.color.rgb = RGBColor(120, 120, 120)


def content_slide(prs, title, bullets, image=None, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT
    add_title(slide, title, subtitle)
    add_bullets(slide, bullets, width=5.35 if image else 11.5)
    if image:
        add_picture(slide, image)
    add_footer(slide, len(prs.slides))
    return slide


def build() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    title = slide.shapes.add_textbox(Inches(0.9), Inches(1.65), Inches(11.5), Inches(1.4))
    p = title.text_frame.paragraphs[0]
    p.text = "CORES for Monocular Depth Estimation"
    p.font.size = Pt(38); p.font.bold = True; p.font.color.rgb = RGBColor(255, 255, 255)
    subtitle = slide.shapes.add_textbox(Inches(0.95), Inches(3.25), Inches(10.8), Inches(1.2))
    s = subtitle.text_frame.paragraphs[0]
    s.text = "Computer Vision — Project 2\nNames and student IDs"
    s.font.size = Pt(22); s.font.color.rgb = RGBColor(220, 228, 240)

    content_slide(prs, "Problem and research question", [
        "Monocular depth models can produce plausible predictions outside their training domain.",
        "Train a lightweight depth estimator on indoor NYU Depth v2.",
        "Treat outdoor KITTI as out-of-distribution.",
        "Question: can signed convolutional responses reliably detect the shift?",
    ])
    content_slide(prs, "Experimental pipeline", [
        "NYU RGB → FastDepth-style MobileNetV2 → signed responses → CORES score.",
        "42,826 train / 4,758 validation / 654 NYU test images.",
        "1,000 KITTI OOD images; common RGB resolution 224 × 304.",
        "Depth: RMSE, AbsRel, δ1–δ3. OOD: AUROC and FPR95.",
    ])
    content_slide(prs, "Depth-estimation baseline", [
        "2.39M parameters; MobileNetV2 encoder and depthwise FastDepth-style decoder.",
        "AdamW, masked metric L1, mixed precision, 20 epochs.",
        "Best checkpoint selected at epoch 19 using NYU validation only.",
        "KITTI OOD: RMSE 9.8854, AbsRel 0.6912, δ1 0.0792 at 0–80 m.",
        "Recovery checkpoints make long Kaggle runs resumable.",
    ], "01_fastdepth_training.png")
    content_slide(prs, "CORES adaptation", [
        "For every kernel: signed spatial maximum and minimum.",
        "Magnitude measures distance beyond calibrated positive/negative thresholds.",
        "Frequency measures how often thresholds are crossed.",
        "Classifier backtracking has no direct dense-regression analogue: evaluate all channels and a 20% extreme-channel variant.",
    ])
    content_slide(prs, "Where is the OOD signal?", [
        "Early encoder: AUROC 0.9897, FPR95 0.044.",
        "Middle encoder: AUROC 0.9999, FPR95 0.001.",
        "Late encoder: AUROC 0.9028, FPR95 0.466.",
        "Late decoder: AUROC 0.5144, FPR95 0.798.",
    ], "04_cores_roc_curves.png")
    content_slide(prs, "Main result: magnitude is sufficient", [
        "Middle encoder, magnitude-only: AUROC 0.999943 and FPR95 0.000.",
        "Response frequency slightly degrades detection.",
        "The decoder is close to random.",
        "Intermediate encoder responses retain the strongest domain information.",
    ], "05_cores_component_ablation.png")
    content_slide(prs, "Controls: not just colour or random features", [
        "Trained FastDepth: 0.99994 / 0.000.",
        "ImageNet encoder without depth training: 0.91886 / 0.324.",
        "Random encoder: 0.45062 / 1.000.",
        "RGB Mahalanobis baseline: 0.82283 / 0.759.",
        "Depth training creates most of the useful response separation.",
    ], "06_cores_robustness_controls.png")
    content_slide(prs, "Multi-layer contribution", [
        "All-layer mean: 0.999737 / 0.001.",
        "Encoder-only mean: 0.999884 / 0.000.",
        "Removing middle encoder: 0.987777 / 0.052.",
        "No aggregation beats middle magnitude-only.",
        "Pure synthetic noise saturates the weighting criterion.",
    ], "07_cores_multilayer_aggregation.png")
    content_slide(prs, "Calibration stability", [
        "Five seeds × five calibration-set sizes.",
        "Only 16 ID images: mean AUROC 0.999943.",
        "FPR95 remains zero in all 25 trials.",
        "Synthetic threshold protocol: AUROC 0.99868, FPR95 0.007.",
    ], "08_cores_calibration_stability.png")
    content_slide(prs, "Architecture ablation: ResNet18", [
        "ResNet18: 11.38M parameters versus MobileNetV2's 2.39M.",
        "NYU-test delta1 improves from 0.7181 to 0.7652.",
        "Middle-magnitude CORES falls from 0.999943 / 0.000 to 0.834523 / 0.499.",
        "ImageNet-only ResNet18 is stronger than its depth-trained version.",
        "Depth quality does not predict OOD quality; the response signal is architecture-dependent.",
    ], "09_architecture_comparison.png")
    content_slide(prs, "Failure cases and limitations", [
        "NYU indoor versus KITTI roads is a broad, visually easy domain shift.",
        "The NYU model is capped at 10 m, so KITTI metric depth deteriorates sharply.",
        "Gaussian near-OOD corruption: AUROC 0.94273, FPR95 0.388.",
        "One trained checkpoint per architecture; parameter count is not isolated.",
        "FastDepth-style MobileNetV2, not an exact MobileNetV1 reproduction.",
        "Dense-prediction channel selection is an explicit adaptation of CORES.",
    ])
    content_slide(prs, "Conclusions", [
        "CORES transfers effectively from classification to dense depth estimation.",
        "The middle encoder magnitude is the simplest and strongest detector.",
        "Weak late/decoder layers can dilute multi-layer aggregation.",
        "The result is stable, but naturally occurring near-OOD remains open.",
        "ResNet18 confirms that near-perfect CORES separation is not universal.",
        "Repository contains notebook, checkpoints, raw tables, figures, tests and report.",
    ])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
